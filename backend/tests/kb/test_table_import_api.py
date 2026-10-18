"""表格文件导入端点 HTTP 冒烟：自动定义列 / 按定义校验 / 不符则拒绝导入。

服务层语义由 ``test_table_schema.py`` 覆盖，这里验证路由接线、结构化 422 detail
（前端按 ``code`` 做多语言提示）以及导入后条目可按文档分页读出。
"""
from __future__ import annotations

from types import SimpleNamespace

import pytest
from fastapi import FastAPI
from fastapi.testclient import TestClient
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.db.database import Base
from app.deps import get_current_user, get_db
from app.models.kb.kb_collection import KbCollection
from app.models.kb.kb_document import KbDocument
from app.models.kb.kb_segment import KbSegment
from app.models.wiki.wiki_knowledge import WikiKnowledge
from app.routers.kb.kb import router

TENANT = 100

CSV_OK = "商品,价格\n防水低价,9\n抗压防水,15\n".encode()
CSV_WRONG_COLS = "商品,库存\nA,3\n".encode()
CSV_WRONG_TYPE = "商品,价格\nA,九块九\n".encode()


@pytest.fixture
def env(db: Session):
    Base.metadata.create_all(db.bind, tables=[KbDocument.__table__])
    kn = WikiKnowledge(tenant_id=TENANT, name="table-kb", slug="table-kb-api",
                       type=2, kb_format="table", index_mode="high_quality")
    db.add(kn)
    db.commit()
    coll = KbCollection(tenant_id=TENANT, name=f"kb_{kn.id}", dimensions=768,
                        knowledge_id=kn.id)
    db.add(coll)
    db.commit()
    # 路由现按 collection_id 调用；暴露 id 供测试构造 URL
    return SimpleNamespace(id=kn.id, collection_id=coll.id)


@pytest.fixture
def client(db: Session):
    app = FastAPI()
    app.include_router(router)
    app.dependency_overrides[get_db] = lambda: db
    app.dependency_overrides[get_current_user] = lambda: SimpleNamespace(
        user_id=1, is_admin=True, tenant_id=TENANT
    )
    with TestClient(app) as c:
        yield c


@pytest.fixture
def inserts(monkeypatch) -> list[dict]:
    """记录每次入库的切片：替掉 ``_insert_via_kb``，避免真实嵌入（不联网）。"""
    calls: list[dict] = []

    def _fake(_session_factory, **kwargs):
        chunk = kwargs["chunks"][0]
        calls.append({
            "n": len(kwargs["chunks"]),
            "document_id": kwargs["document_id"],
            "meta": dict(chunk.metadata or {}),
            "content": chunk.content.text,
        })

    monkeypatch.setattr("app.services.kb.rag.content_service._insert_via_kb", _fake)
    return calls


def _import_csv(client: TestClient, env, data: bytes, name: str = "t.csv"):
    return client.post(
        f"/api/v1/kb/collections/{env.collection_id}/table-records/import",
        files={"file": (name, data, "text/csv")},
    )


def test_parse_table_picks_data_sheet_when_active_is_cover():
    """多 sheet 工作簿：激活页是封面（无数据行）时，parse_table 应取含数据的页，
    而非误报「表格文件没有数据行」。"""
    import io

    import openpyxl

    from app.services.kb.rag.content_service import parse_table

    wb = openpyxl.Workbook()
    cover = wb.active
    cover.title = "封面"
    cover.append(["特斯拉 Optimus 供应商总表"])  # 仅标题，无数据行
    data = wb.create_sheet("供应商")
    data.append(["公司", "代码", "备注"])
    data.append(["A 公司", "600000", "电机"])
    data.append(["B 公司", "000001", "减速器"])

    buf = io.BytesIO()
    wb.save(buf)
    raw = buf.getvalue()

    columns, rows = parse_table(raw, "总表.xlsx")
    assert columns == ["公司", "代码", "备注"]
    assert len(rows) == 2
    assert rows[0]["公司"] == "A 公司"
    assert rows[1]["代码"] == "000001"


def test_import_auto_creates_schema_and_registers_document(db, env, client, inserts):
    r = _import_csv(client, env, CSV_OK, "价格表.csv")
    assert r.status_code == 201, r.text
    body = r.json()
    assert body["ingested"] == 2 and body["schema_created"] is True
    # 自动生成字段定义时默认**所有字段**组合为向量化内容
    assert body["embed_fields"] == ["商品", "价格"]
    assert body["embed_field"] == "商品"
    assert [f["name"] for f in body["schema"]] == ["商品", "价格"]
    assert next(f for f in body["schema"] if f["name"] == "价格")["type"] == "number"
    # 切片 document_id 必须等于文档行 uuid_code，「表格条目」按文档取切片才拿得到
    assert inserts[0]["document_id"] == body["document_id"]
    # 一行一段：content 是向量化字段组合拼成的 CSV 行（逗号间隔）
    assert inserts[0]["content"] == "防水低价,9"
    assert inserts[0]["meta"]["价格"] == 9

    doc = db.execute(
        select(KbDocument).where(KbDocument.uuid_code == body["document_id"])
    ).scalar_one()
    assert doc.knowledge_id == env.id and doc.name == "价格表.csv"
    assert doc.segment_count == 2 and doc.file_type == "csv"

    schema = client.get(f"/api/v1/kb/collections/{env.collection_id}/schema").json()
    assert schema["embed_field"] == "商品"
    assert schema["embed_fields"] == ["商品", "价格"]
    assert schema["columns"] == ["商品", "价格"]
    # 文档必有所属集合（多集合场景按 collection_id 精确归属）
    assert doc.collection_id is not None


def test_import_validates_existing_schema_and_rejects_mismatch(db, env, client, inserts):
    assert _import_csv(client, env, CSV_OK).status_code == 201
    before = len(inserts)

    missing = _import_csv(client, env, CSV_WRONG_COLS, "b.csv")
    assert missing.status_code == 422
    detail = missing.json()["detail"]
    assert detail["code"] == "table_schema_mismatch"
    assert detail["missing"] == ["价格"] and detail["extra"] == ["库存"]

    bad_type = _import_csv(client, env, CSV_WRONG_TYPE, "c.csv")
    assert bad_type.status_code == 422
    cell = bad_type.json()["detail"]
    assert cell["code"] == "table_cell_type_mismatch"
    assert cell["column"] == "价格" and cell["expected"] == "number"

    # 两次拒绝都不落数据：只有首次导入产生过切片
    assert len(inserts) == before


def test_import_accepts_different_column_order(db, env, client, inserts):
    assert _import_csv(client, env, CSV_OK).status_code == 201
    reordered = _import_csv(client, env, "价格,商品\n20,B\n".encode())
    assert reordered.status_code == 201, reordered.text
    assert reordered.json()["embed_field"] == "商品"
    assert inserts[-1]["meta"]["价格"] == 20


def test_import_rejects_unsupported_file_type(client, env, inserts):
    r = _import_csv(client, env, b"%PDF-1.4", "a.pdf")
    assert r.status_code == 422
    assert not inserts


def test_preview_validates_without_persisting(db, env, client, inserts):
    url = f"/api/v1/kb/collections/{env.collection_id}/table-records/preview"
    # 尚无定义：预览只做规范化推导，不落库也不产生切片
    ok = client.post(url, files={"file": ("t.csv", CSV_OK, "text/csv")})
    assert ok.status_code == 200
    assert ok.json()["columns"] == ["商品", "价格"]
    assert ok.json()["embed_field"] == "商品"
    assert client.get(f"/api/v1/kb/collections/{env.collection_id}/schema").json()["fields"] == []
    assert not inserts

    # 已有定义后，预览与导入同口径：不匹配同样 422
    assert _import_csv(client, env, CSV_OK).status_code == 201
    bad = client.post(url, files={"file": ("t.csv", CSV_WRONG_TYPE, "text/csv")})
    assert bad.status_code == 422
    assert bad.json()["detail"]["code"] == "table_cell_type_mismatch"


def test_put_schema_requires_single_embedded_column(client, env):
    bad = client.put(
        f"/api/v1/kb/collections/{env.collection_id}/schema",
        json={"fields": [{"name": "a", "type": "string", "embedding": False}]},
    )
    assert bad.status_code == 422
    assert bad.json()["detail"]["code"] == "table_schema_embed_required"

    ok = client.put(
        f"/api/v1/kb/collections/{env.collection_id}/schema",
        json={"fields": [
            {"name": "a", "type": "string", "embedding": True},
            {"name": "b", "type": "number", "embedding": False},
        ]},
    )
    assert ok.status_code == 200 and ok.json()["embed_field"] == "a"


def test_document_segments_paginate_and_filter_by_keyword(db, env, client):
    db.add(KbDocument(tenant_id=TENANT, uuid_code="doc-t", knowledge_id=env.id,
                      collection=f"kb_{env.id}", name="t.csv", source_type="db_table",
                      status="completed"))
    db.add_all([
        KbSegment(tenant_id=TENANT, collection=f"kb_{env.id}", document_id="doc-t",
                  chunk_index=i, content=c, chunk_type="table_row",
                  metadata_={"chunk_type": "table_row"})
        for i, c in enumerate(["防水低价", "抗压支架"])
    ])
    db.commit()

    page = client.get(f"/api/v1/kb/documents/doc-t/segments",
                      params={"chunk_type": "table_row", "page": 1, "page_size": 1})
    assert page.status_code == 200
    body = page.json()
    assert body["total"] == 2 and len(body["items"]) == 1

    hit = client.get(f"/api/v1/kb/documents/doc-t/segments",
                     params={"chunk_type": "table_row", "q": "防水"})
    assert hit.json()["total"] == 1
    assert hit.json()["items"][0]["content"] == "防水低价"


def test_retry_table_file_route_accepts_persisted_file(db, env, client, inserts):
    """表格文件重试路由：导入时已落盘原始文件，重试不被 400 拒绝（走表格分支）。"""
    r = _import_csv(client, env, CSV_OK, "重试表.csv")
    assert r.status_code == 201, r.text
    doc_id = r.json()["document_id"]
    doc = db.execute(
        select(KbDocument).where(KbDocument.uuid_code == doc_id)
    ).scalar_one()
    # 导入时已落盘原始文件 → meta.file_path 存在，重试不会被 400 拒绝
    assert (doc.meta or {}).get("file_path")

    retry = client.post(f"/api/v1/kb/documents/{doc_id}/retry")
    assert retry.status_code == 200, retry.text


def test_import_table_reuses_document_id(db, env, inserts):
    """import_table_file 用同一 document_id + create_document=False 重导不产生重复文档行。"""
    from app.services.kb.rag.content_service import import_table_file

    res1 = import_table_file(db, TENANT, f"kb_{env.id}", CSV_OK, "a.csv",
                             knowledge_id=env.id)
    doc_id = res1["document_id"]
    assert res1["ingested"] == 2

    # 以同一文档行重导（模拟「重试」），不新建 kms_document
    res2 = import_table_file(
        db, TENANT, f"kb_{env.id}", CSV_OK, "a.csv",
        knowledge_id=env.id, document_id=doc_id, create_document=False,
    )
    assert res2["ingested"] == 2
    assert res2["document_id"] == doc_id

    docs = db.execute(
        select(KbDocument).where(KbDocument.collection == f"kb_{env.id}")
    ).scalars().all()
    assert len(docs) == 1
    assert docs[0].segment_count == 2


def test_retry_table_file_without_persisted_file_is_rejected(db, env, client):
    """未留存原始文件的表格文档（历史数据）重试应 400 提示改用重新向量化。"""
    doc = KbDocument(tenant_id=TENANT, uuid_code="no-file", knowledge_id=env.id,
                     collection=f"kb_{env.id}", name="历史.csv", source_type="db_table",
                     status="completed", meta={"source": "table-import"})
    db.add(doc)
    db.commit()

    r = client.post("/api/v1/kb/documents/no-file/retry")
    assert r.status_code == 400
    assert "重新向量化" in r.json()["detail"]