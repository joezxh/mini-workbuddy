"""QA / 表格行 / 多模态资产的内容服务（对齐文档 §8 / D15 / D9）。

共同形态：**不经 parser/chunker**，直接构造 ``Chunk`` 调
``KnowledgeBase.insert_document``（嵌入在 KB 内部完成）。
``kb_factory`` 可注入（测试替换真实嵌入模型）；生产路径传 ``session_factory``。

D9 边界：AgentScope 2.0.8 ``supports_multimodal`` 恒 False → 仅支持**文搜图**
（图片以 caption/OCR 文本入库 + 资产文件落盘），图搜图不支持。
"""
from __future__ import annotations

import asyncio
import csv
import io
import uuid
from pathlib import Path
from typing import Callable, Optional

from sqlalchemy import text
from sqlalchemy.orm import Session

from app.models.kb.kb_document import KbDocument, new_document_uuid
from app.models.kb.kb_segment import KbSegment
from app.models.kb.kb_segment_asset import KbSegmentAsset
from app.services.kb.rag.chunker_factory import build_qa_chunks, build_table_row_chunks

ASSET_MAX_BYTES = 2 * 1024 * 1024   # 单图 ≤ 2MB（对齐 Dify）
ASSET_MAX_PER_SEGMENT = 10          # 单 chunk ≤ 10 张


def _insert_via_kb(
    session_factory: Optional[Callable],
    *,
    db: Session,
    collection: str,
    tenant_id: int,
    document_id: str,
    chunks,
    kb_factory: Optional[Callable] = None,
    document_metadata: Optional[dict] = None,
) -> None:
    """经 KnowledgeBase.insert_document 落库（嵌入内部完成）。

    economy 模式写 NULL 向量（D11）；检索走独立关键词服务（AC7）。

    本函数是同步函数，但内部要跑异步 KB 写入。调用方既有后台线程（无运行中的
    事件循环），也有 async 路由（已在事件循环中）。``asyncio.run()`` 不能在已运行
    的事件循环里调用，故：有运行循环时，把协程丢到独立线程里用 ``asyncio.run``
    跑（新线程自有事件循环）；否则直接跑。``db`` 的读取（resolve_index_mode）
    留在调用线程，避免跨线程使用 SQLAlchemy Session。
    """
    from app.services.kb.index_mode import resolve_index_mode

    # 在调用线程先解析索引模式（仅一次 DB 读，避免跨线程用 Session）
    index_mode = resolve_index_mode(db, collection)

    async def _run() -> None:
        if kb_factory is not None:
            async with kb_factory(index_mode=index_mode) as kb:
                await kb.ensure_collection()
                await kb.insert_document(chunks, document_id=document_id,
                                         document_metadata=document_metadata)
            return
        from app.db.database import SessionLocal
        from app.services.kb.rag.knowledge_factory import knowledge_base

        async with knowledge_base(
            SessionLocal, collection_name=collection, tenant_id=tenant_id,
            name=collection, index_mode=index_mode,
        ) as kb:
            await kb.ensure_collection()
            await kb.insert_document(chunks, document_id=document_id,
                                     document_metadata=document_metadata)

    try:
        asyncio.get_running_loop()
    except RuntimeError:
        # 无运行中的事件循环（如后台线程）：直接跑
        asyncio.run(_run())
        return

    # 已在事件循环中（如 async 路由）：独立线程里跑 asyncio.run
    import concurrent.futures

    with concurrent.futures.ThreadPoolExecutor(max_workers=1) as pool:
        pool.submit(asyncio.run, _run()).result()


def _track_document(
    db: Session, tenant_id: int, collection: str, name: str,
    source_type: str, count: int,
    *,
    uuid_code: Optional[str] = None,
    knowledge_id: Optional[int] = None,
    file_type: Optional[str] = None,
    file_size: Optional[int] = None,
    meta: Optional[dict] = None,
    collection_id: Optional[int] = None,
) -> KbDocument:
    """登记一条文档行。

    ``uuid_code`` 必须与切片里的 ``document_id`` 一致（``kb_segment.document_id``
    的值即本表 ``uuid_code``），否则文档列表里的条目点开看不到任何切片。
    ``collection_id`` 显式记录所属集合（表格知识库多集合场景下文档必选集合）。
    """
    doc = KbDocument(
        tenant_id=tenant_id,
        uuid_code=uuid_code or new_document_uuid(),
        knowledge_id=knowledge_id if knowledge_id is not None else 0,
        collection=collection,
        collection_id=collection_id,
        name=name,
        source_type=source_type,
        status="completed",
        segment_count=count,
        file_type=file_type,
        file_size=file_size,
        meta=meta,
    )
    db.add(doc)
    db.commit()
    db.refresh(doc)
    return doc


# ── Q&A（D15）────────────────────────────────────────────────────────────

def ingest_qa_records(
    db: Session, tenant_id: int, collection: str, records: list[dict],
    session_factory: Optional[Callable] = None,
    kb_factory: Optional[Callable] = None,
) -> int:
    """Q&A 直构 Chunk：content=问题（仅问题被嵌入），answer/tags 进 metadata。"""
    if not records:
        return 0
    doc_id = new_document_uuid()
    chunks = build_qa_chunks(records, source="qa-import")
    _insert_via_kb(
        session_factory, db=db, collection=collection, tenant_id=tenant_id,
        document_id=doc_id, chunks=chunks, kb_factory=kb_factory,
        document_metadata={"source": "qa-import"},
    )
    _track_document(db, tenant_id, collection, f"qa-{doc_id[:8]}", "qa_import", len(chunks))
    return len(chunks)


def list_qa_records(db: Session, tenant_id: int, collection: str,
                    limit: int = 100) -> list[dict]:
    """Q&A 条目列表；``segment_id`` 供前端直接改单条（启停/编辑）而不必回表定位。"""
    rows = db.execute(text(
        "SELECT id, document_id, chunk_index, content, metadata"
        " FROM kms_segment WHERE collection = :c AND tenant_id = :t"
        " AND metadata->>'chunk_type' = 'qa' ORDER BY id DESC LIMIT :k"
    ), {"c": collection, "t": tenant_id, "k": limit}).all()
    return [
        {
            "segment_id": r.id,
            "document_id": r.document_id,
            "chunk_index": r.chunk_index,
            "question": r.content,
            "answer": (r.metadata or {}).get("answer"),
            "tags": (r.metadata or {}).get("tags") or [],
            "enabled": (r.metadata or {}).get("enabled", True),
        }
        for r in rows
    ]


# ── 表格行（D15）─────────────────────────────────────────────────────────

def parse_table(data: bytes, filename: str, limit: int | None = None) -> tuple[list[str], list[dict]]:
    """CSV / xlsx → ``(表头列名, [{列名: 值}])``。

    列名去首尾空白并丢弃空列名；``limit`` 限制返回的**数据行**数（表头恒返回）。
    xlsx 读出的 ``datetime`` 等对象原样交给上层按列定义规范化。
    """
    rows: list[dict] = []
    lower = filename.lower()
    if lower.endswith(".csv"):
        reader = csv.DictReader(io.StringIO(data.decode("utf-8-sig")))
        headers = [(h or "").strip() for h in (reader.fieldnames or [])]
        columns = [h for h in headers if h]
        for raw in reader:
            row = {
                str(k).strip(): v for k, v in raw.items()
                if k and str(k).strip() and str(k).strip() in columns
            }
            rows.append(row)
            if limit and len(rows) >= limit:
                break
        return columns, rows
    if lower.endswith((".xlsx", ".xls")):
        import openpyxl

        def _read_workbook(wb) -> tuple[list[str], list[dict]]:
            # 优先用含数据最多的工作表；若为封面/空页导致 0 行，再回退尝试其余工作表
            candidate = _pick_data_sheet(wb)
            sheets_to_try = [candidate] + [s for s in wb.worksheets if s is not candidate]
            columns: list[str] = []
            rows: list[dict] = []
            for ws in sheets_to_try:
                iterator = ws.iter_rows(values_only=True)
                headers = [str(h).strip() if h is not None else "" for h in next(iterator, [])]
                columns = [h for h in headers if h]
                if not columns:
                    continue
                rows = []
                for values in iterator:
                    if all(v is None for v in values):
                        continue
                    row = {
                        h: v for h, v in zip(columns, values) if h
                    }
                    if not any(not _cell_empty(v) for v in row.values()):
                        continue
                    rows.append(row)
                    if limit and len(rows) >= limit:
                        break
                if rows:
                    break
            return columns, rows

        # 先以只读模式解析（省内存、快）。已知 openpyxl 在 read_only=True 配合
        # data_only=True 时，部分 xlsx 读不到单元格缓存值（全返回 None）而致 0 行；
        # 此时回退用非只读模式重读，确保 data_only 能取到真实值。若仍 0 行，最后
        # 以 data_only=False 兜底（非公式单元格值不变，仅在确为公式无缓存时退化
        # 为公式串），避免直接报「没有数据行」。
        columns: list[str] = []
        rows = []
        wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        try:
            columns, rows = _read_workbook(wb)
        finally:
            wb.close()
        if not rows:
            wb2 = openpyxl.load_workbook(io.BytesIO(data), read_only=False, data_only=True)
            try:
                columns, rows = _read_workbook(wb2)
            finally:
                wb2.close()
        if not rows:
            wb3 = openpyxl.load_workbook(io.BytesIO(data), read_only=False, data_only=False)
            try:
                columns, rows = _read_workbook(wb3)
            finally:
                wb3.close()
        return columns, rows
    raise ValueError(f"不支持的表格文件类型: {filename!r}（仅 csv/xlsx/xls）")


def _cell_empty(value) -> bool:
    return value is None or (isinstance(value, str) and not value.strip())


def _pick_data_sheet(wb) -> "object":
    """多 sheet 工作簿里挑出最可能含数据的工作表。

    优先用激活页；若其行数 ≤ 1（封面 / 汇总 / 空页），改选数据行最多的工作表。
    这样可避免「激活页非数据页」导致解析到 0 行而误报「表格文件没有数据行」。
    """
    sheets = list(wb.worksheets)
    if not sheets:
        return wb.active
    active = wb.active
    if active is None:
        active = sheets[0]

    def _rowcount(ws) -> int:
        try:
            return int(ws.max_row or 0)
        except Exception:
            return 0

    if _rowcount(active) > 1:
        return active
    best = max(sheets, key=_rowcount)
    return best if _rowcount(best) > 1 else active


def parse_table_bytes(data: bytes, filename: str, limit: int | None = None) -> list[dict]:
    """CSV / xlsx → [{列名: 值}]（仅取数据行，表头经 :func:`parse_table` 获取）。"""
    return parse_table(data, filename, limit=limit)[1]


def normalize_embed_fields(value) -> list[str]:
    """把「向量化字段组合」入参收敛为字段名列表（兼容旧的单个字段名 ``str``）。"""
    if isinstance(value, str):
        return [value] if value else []
    return [str(v) for v in (value or []) if v]


def ingest_table_rows(
    db: Session, tenant_id: int, collection: str, rows: list[dict],
    embed_fields,
    session_factory: Optional[Callable] = None,
    kb_factory: Optional[Callable] = None,
    document_id: Optional[str] = None,
    *,
    create_document: bool = True,
    knowledge_id: Optional[int] = None,
    doc_name: Optional[str] = None,
    file_type: Optional[str] = None,
    file_size: Optional[int] = None,
    file_path: Optional[str] = None,
    collection_id: Optional[int] = None,
) -> int:
    """表格行直构 Chunk：**一行一段**，内容为向量化字段组合拼成的 CSV 行。

    ``embed_fields`` 是集合配置的向量化字段组合（1~N 个字段名；传单个 ``str`` 时
    按单字段处理以兼容旧调用）。每段：
    * ``content`` = 这些字段值拼成的 CSV 行（逗号间隔，RFC4180 转义）——即向量化内容；
    * ``metadata`` = 该行全部字段，供按集合字段定义回显记录与过滤。

    ``document_id`` 指定切片用的文档 ID（缺省现生成）；``create_document=False`` 时
    不新建文档行——定时同步覆盖同一文档（spec §10.8）用它避免重复文档行。
    新建文档行时 ``uuid_code`` 与切片的 ``document_id`` 取同一值，否则「表格条目」
    按文档取切片会取不到（``kb_segment.document_id`` 的值即 ``kms_document.uuid_code``）。
    ``file_path`` 为落盘的原始文件路径（重试重导用），仅新建文档行时写入 meta。
    ``collection_id`` 显式记录文档所属集合（表格库多集合场景必选）。
    """
    if not rows:
        return 0
    names = normalize_embed_fields(embed_fields)
    if not names:
        raise ValueError("未指定向量化字段组合（embed_fields）")
    missing = [name for name in names if name not in rows[0]]
    if missing:
        raise ValueError(f"embedding 字段 {missing} 不在表头中")
    doc_id = document_id or new_document_uuid()
    chunks = build_table_row_chunks(rows, embed_fields=names, source="table-import")
    _insert_via_kb(
        session_factory, db=db, collection=collection, tenant_id=tenant_id,
        document_id=doc_id, chunks=chunks, kb_factory=kb_factory,
        document_metadata={"source": "table-import", "embed_fields": names},
    )
    if create_document:
        doc_meta = {"source": "table-import", "embed_fields": names}
        if file_path:
            doc_meta["file_path"] = file_path
        _track_document(
            db, tenant_id, collection, doc_name or f"table-{doc_id[:8]}",
            "db_table", len(chunks),
            uuid_code=doc_id, knowledge_id=knowledge_id,
            file_type=file_type, file_size=file_size,
            meta=doc_meta, collection_id=collection_id,
        )
    return len(chunks)


def import_table_file(
    db: Session,
    tenant_id: int,
    collection: str,
    data: bytes,
    filename: str,
    *,
    knowledge_id: Optional[int] = None,
    embed_field: Optional[str] = None,
    embed_fields: Optional[list[str]] = None,
    session_factory: Optional[Callable] = None,
    kb_factory: Optional[Callable] = None,
    document_id: Optional[str] = None,
    create_document: bool = True,
    collection_id: Optional[int] = None,
    auto_create_collection: bool = False,
) -> dict:
    """表格文件导入：解析 → **列定义自动生成/校验** → 行级入库 + 文档登记。

    * 知识库没有列定义 → 按文件表头自动定义并落库（**默认全部字段**为向量化字段）；
    * 已有列定义 → 缺列/多列/取值类型不符抛
      :class:`~app.services.kb.table_schema.TableSchemaMismatch`，**不落任何数据**。

    集合归属（表格知识库支持一个知识库多个集合）：
    * ``collection_id`` 指定 → 取该集合并校验其属于 ``knowledge_id``；
    * 未指定集合且 ``auto_create_collection`` → 在该知识库下**自动新建**集合
      （名称 ``kb_{kid}_{slug}``，slug 取文件名），承载本次推断出的字段定义；
    * 仅给 ``collection`` 名 → 沿用既有集合行（缺失时补建），行为不变。

    ``document_id`` / ``create_document=False`` 用于「重试重导」：复用同一文档行
    （``uuid_code`` 不变），不新建 ``kms_document``。导入时把原始文件落盘到
    ``kb_document.meta['file_path']``，使后续「重试」能读回字节重跑。

    返回 ``{ingested, rows, document_id, embed_fields, embed_field, columns, schema,
    schema_created, collection, collection_id}``。
    """
    from app.services.kb.collection_service import (
        create_knowledge_collection,
        get_collection_by_id,
        validate_collection_ownership,
    )
    from app.services.kb.table_schema import ensure_collection_row, resolve_import_schema
    from app.services.kb.upload_store import save_upload

    columns, rows = parse_table(data, filename)
    if not columns:
        raise ValueError("表格文件没有可用表头（首行需为列名）")
    if not rows:
        raise ValueError("表格文件没有数据行")

    # ── 解析目标集合（保证文档必有所属集合，并拿到 collection_id）──
    coll = None
    if collection_id:
        coll = get_collection_by_id(db, collection_id)
        if coll is None:
            raise ValueError(f"集合不存在: collection_id={collection_id}")
    elif auto_create_collection and not collection:
        if knowledge_id is None:
            raise ValueError("自动新建集合需要 knowledge_id")
        coll = create_knowledge_collection(db, tenant_id, knowledge_id, slug=filename)
    else:
        # 既有路径：按名取集合行，缺失则补建（HTTP 建库不建集合行的兜底）
        coll = ensure_collection_row(db, collection, tenant_id=tenant_id)
    validate_collection_ownership(db, coll, knowledge_id)
    collection = coll.name
    collection_id = coll.id

    schema = resolve_import_schema(
        db, collection, columns, rows,
        embed_field=embed_field, embed_fields=embed_fields, tenant_id=tenant_id,
    )
    targets = schema["embed_fields"]
    doc_id = document_id or new_document_uuid()
    # 落盘原始文件（失败不阻塞导入，仅告警；此时重试会回退到「重新向量化」提示）
    file_path = save_upload(tenant_id, doc_id, filename, data)
    ingested = ingest_table_rows(
        db, tenant_id, collection, schema["rows"], targets,
        session_factory=session_factory, kb_factory=kb_factory,
        document_id=doc_id, create_document=create_document,
        knowledge_id=knowledge_id, doc_name=filename,
        file_type=(filename or "").rsplit(".", 1)[-1].lower() or None,
        file_size=len(data), file_path=file_path,
        collection_id=collection_id,
    )
    return {
        "ingested": ingested,
        "rows": len(rows),
        "document_id": doc_id,
        # 向量化字段组合（1~N 个）；embed_field 为其中首个，兼容旧调用方
        "embed_fields": targets,
        "embed_field": targets[0] if targets else None,
        "collection": collection,
        "collection_id": collection_id,
        "columns": columns,
        "schema": schema["fields"],
        "schema_created": schema["created"],
    }


# ── 多模态资产（D9：仅文搜图）────────────────────────────────────────────

def save_image_asset(
    db: Session, tenant_id: int, collection: str,
    data: bytes, filename: str, mime_type: str, caption: str,
    storage_dir: Path,
    session_factory: Optional[Callable] = None,
    kb_factory: Optional[Callable] = None,
) -> KbSegmentAsset:
    """图片落盘 + caption 文本 chunk 入库（文搜图）+ 资产行登记。

    资产行挂在 insert_document 产生的首个 kms_segment 行上（D9：原生 Chunk 无
    segment id，取该文档 chunk_index=0 的行）。
    """
    if len(data) > ASSET_MAX_BYTES:
        raise ValueError(f"图片超过 2MB 上限（{len(data)} 字节）")
    if not mime_type.startswith("image/"):
        raise ValueError(f"仅支持图片类型，收到 {mime_type!r}")

    doc_id = new_document_uuid()
    from agentscope.message import TextBlock
    from agentscope.rag import Chunk

    chunk = Chunk(
        content=TextBlock(text=caption or filename),
        source=filename,
        chunk_index=0, total_chunks=1,
        metadata={"chunk_type": "image", "filename": filename, "mime": mime_type},
    )
    _insert_via_kb(
        session_factory, db=db, collection=collection, tenant_id=tenant_id,
        document_id=doc_id, chunks=[chunk], kb_factory=kb_factory,
        document_metadata={"source": "image", "filename": filename},
    )

    storage_dir = Path(storage_dir)
    storage_dir.mkdir(parents=True, exist_ok=True)
    asset_name = f"{uuid.uuid4().hex}_{filename}"
    path = storage_dir / asset_name
    path.write_bytes(data)

    segment = db.execute(
        text("SELECT id FROM kms_segment WHERE collection=:c AND document_id=:d"
             " AND chunk_index=0 LIMIT 1"),
        {"c": collection, "d": doc_id},
    ).scalar_one()

    asset = KbSegmentAsset(
        tenant_id=tenant_id, segment_id=segment,
        file_path=str(path), mime_type=mime_type, size=len(data),
    )
    db.add(asset)
    db.commit()
    db.refresh(asset)
    return asset


def list_assets(db: Session, tenant_id: int, collection: str,
                limit: int = 100) -> list[dict]:
    """按 collection 列出图片资产（经 kms_segment 反查）。"""
    rows = db.execute(text(
        "SELECT a.id, a.file_path, a.mime_type, a.size, a.created_at"
        " FROM kms_segment_asset a"
        " JOIN kms_segment s ON s.id = a.segment_id"
        " WHERE s.collection = :c AND a.tenant_id = :t"
        " ORDER BY a.id DESC LIMIT :k"
    ), {"c": collection, "t": tenant_id, "k": limit}).all()
    return [
        {"id": r.id, "file_path": r.file_path, "mime_type": r.mime_type,
         "size": r.size, "created_at": str(r.created_at) if r.created_at else None}
        for r in rows
    ]
