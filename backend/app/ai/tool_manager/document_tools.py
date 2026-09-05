"""文档处理工具集 - PDF 解析 / Excel 处理 / 文档生成 / 格式转换。

基于 AgentScope 2.0.4 ToolBase 协议实现。
"""
from __future__ import annotations

import json
import logging
import os
from pathlib import Path
from typing import Any, AsyncGenerator, List, Optional

from agentscope.tool import ToolBase, ToolChunk
from agentscope.message import TextBlock
from agentscope.permission import (
    PermissionBehavior,
    PermissionContext,
    PermissionDecision,
)

logger = logging.getLogger(__name__)


def _text_chunk(payload: Any) -> ToolChunk:
    """把任意 python 对象序列化成 TextBlock ToolChunk。"""
    if isinstance(payload, str):
        text = payload
    else:
        text = json.dumps(payload, ensure_ascii=False, default=str)
    return ToolChunk(content=[TextBlock(type="text", text=text)])


# ============= PDFParser =============

class PDFParser(ToolBase):
    """PDF 文件解析 - 提取文本、表格和元数据。"""

    name: str = "pdf_parser"
    description: str = "解析 PDF 文件，提取文本内容、表格数据和元数据信息"
    input_schema: dict = {
        "type": "object",
        "properties": {
            "file_path": {"type": "string", "description": "PDF 文件路径"},
            "pages": {
                "type": "array",
                "items": {"type": "integer"},
                "description": "指定页码列表（从1开始），不传则解析全部",
            },
            "extract_tables": {
                "type": "boolean",
                "description": "是否提取表格",
                "default": True,
            },
        },
        "required": ["file_path"],
    }
    is_concurrency_safe: bool = True
    is_read_only: bool = True

    async def check_permissions(
        self, tool_input: dict, context: PermissionContext,
    ) -> PermissionDecision:
        return PermissionDecision(
            behavior=PermissionBehavior.ALLOW,
            message="PDF parsing is read-only.",
        )

    async def call(
        self,
        file_path: str,
        pages: Optional[List[int]] = None,
        extract_tables: bool = True,
    ) -> ToolChunk:
        try:
            import pdfplumber
        except ImportError:
            return _text_chunk({"error": "pdfplumber 未安装，请执行: pip install pdfplumber"})

        if not os.path.isfile(file_path):
            return _text_chunk({"error": f"文件不存在: {file_path}"})

        try:
            result_pages = []
            with pdfplumber.open(file_path) as pdf:
                total_pages = len(pdf.pages)
                metadata = {
                    "total_pages": total_pages,
                    "metadata": dict(pdf.metadata) if pdf.metadata else {},
                }
                target_pages = pages if pages else range(1, total_pages + 1)
                for page_no in target_pages:
                    if page_no < 1 or page_no > total_pages:
                        continue
                    page = pdf.pages[page_no - 1]
                    page_data = {
                        "page_no": page_no,
                        "text": page.extract_text() or "",
                    }
                    if extract_tables:
                        tables = page.extract_tables() or []
                        page_data["tables"] = tables
                    result_pages.append(page_data)

            return _text_chunk({
                **metadata,
                "pages": result_pages,
            })
        except Exception as exc:
            logger.warning("pdf_parser failed: %s", exc)
            return _text_chunk({"error": str(exc)})


# ============= ExcelProcessor =============

class ExcelProcessor(ToolBase):
    """Excel 文件读写处理。"""

    name: str = "excel_processor"
    description: str = "读取、写入 Excel 文件，支持 sheet 操作、范围读取和数据统计"
    input_schema: dict = {
        "type": "object",
        "properties": {
            "mode": {
                "type": "string",
                "enum": ["read", "write", "info"],
                "description": "操作模式：read-读取 / write-写入 / info-查看信息",
            },
            "file_path": {"type": "string", "description": "Excel 文件路径"},
            "sheet_name": {
                "type": "string",
                "description": "Sheet 名称（不传则操作第一个 sheet）",
            },
            "range": {
                "type": "string",
                "description": "读取范围，如 A1:D10（read 模式）",
            },
            "data": {
                "type": "array",
                "items": {"type": "array", "items": {"type": "string"}},
                "description": "写入数据，二维数组（write 模式）",
            },
            "start_cell": {
                "type": "string",
                "description": "写入起始单元格",
                "default": "A1",
            },
        },
        "required": ["mode", "file_path"],
    }
    is_concurrency_safe: bool = True
    is_read_only: bool = False

    async def check_permissions(
        self, tool_input: dict, context: PermissionContext,
    ) -> PermissionDecision:
        mode = tool_input.get("mode", "read")
        if mode == "read" or mode == "info":
            return PermissionDecision(
                behavior=PermissionBehavior.ALLOW,
                message="Excel read/info is safe.",
            )
        return PermissionDecision(
            behavior=PermissionBehavior.ASK,
            message="Excel write mode will modify the file.",
        )

    async def call(
        self,
        mode: str,
        file_path: str,
        sheet_name: Optional[str] = None,
        range: Optional[str] = None,
        data: Optional[List[List]] = None,
        start_cell: str = "A1",
    ) -> ToolChunk:
        try:
            import openpyxl
        except ImportError:
            return _text_chunk({"error": "openpyxl 未安装，请执行: pip install openpyxl"})

        try:
            if mode == "info":
                if not os.path.isfile(file_path):
                    return _text_chunk({"error": f"文件不存在: {file_path}"})
                wb = openpyxl.load_workbook(file_path, read_only=True, data_only=True)
                info = {
                    "mode": "info",
                    "file_path": file_path,
                    "sheets": [],
                }
                for name in wb.sheetnames:
                    ws = wb[name]
                    info["sheets"].append({
                        "name": name,
                        "dimensions": ws.dimensions,
                        "max_row": ws.max_row,
                        "max_column": ws.max_column,
                    })
                wb.close()
                return _text_chunk(info)

            elif mode == "read":
                if not os.path.isfile(file_path):
                    return _text_chunk({"error": f"文件不存在: {file_path}"})
                wb = openpyxl.load_workbook(file_path, read_only=True, data_only=True)
                ws = wb[sheet_name] if sheet_name and sheet_name in wb.sheetnames else wb.active
                if range:
                    rows = []
                    for row in ws[range]:
                        rows.append([str(cell.value) if cell.value is not None else "" for cell in row])
                else:
                    rows = []
                    for row in ws.iter_rows():
                        rows.append([str(cell.value) if cell.value is not None else "" for cell in row])
                result = {
                    "mode": "read",
                    "file_path": file_path,
                    "sheet_name": ws.title,
                    "dimensions": ws.dimensions,
                    "data": rows,
                    "row_count": len(rows),
                }
                wb.close()
                return _text_chunk(result)

            elif mode == "write":
                if data is None:
                    return _text_chunk({"error": "write 模式必须提供 data 参数"})
                if os.path.isfile(file_path):
                    wb = openpyxl.load_workbook(file_path)
                else:
                    wb = openpyxl.Workbook()
                ws = wb[sheet_name] if sheet_name and sheet_name in wb.sheetnames else (wb[sheet_name] if sheet_name else wb.active)
                if sheet_name and sheet_name not in wb.sheetnames:
                    ws = wb.create_sheet(sheet_name)
                from openpyxl.utils import coordinate_to_tuple
                start_row, start_col = coordinate_to_tuple(start_cell)
                for r_idx, row_data in enumerate(data):
                    for c_idx, value in enumerate(row_data):
                        ws.cell(row=start_row + r_idx, column=start_col + c_idx, value=value)
                wb.save(file_path)
                wb.close()
                return _text_chunk({
                    "mode": "write",
                    "file_path": file_path,
                    "sheet_name": ws.title,
                    "rows_written": len(data),
                    "start_cell": start_cell,
                })
            else:
                return _text_chunk({"error": f"未知 mode: {mode}，支持 read/write/info"})

        except Exception as exc:
            logger.warning("excel_processor failed: %s", exc)
            return _text_chunk({"error": str(exc)})


# ============= DocumentGenerator =============

class DocumentGenerator(ToolBase):
    """文档生成 - 生成 Word / PDF / PPT 文档。"""

    name: str = "document_generator"
    description: str = "生成 Word(.docx)、PDF 或 PPT(.pptx) 文档，支持标题和多段落内容"
    input_schema: dict = {
        "type": "object",
        "properties": {
            "format": {
                "type": "string",
                "enum": ["docx", "pdf", "pptx"],
                "description": "输出文档格式",
            },
            "title": {"type": "string", "description": "文档标题"},
            "sections": {
                "type": "array",
                "items": {
                    "type": "object",
                    "properties": {
                        "heading": {"type": "string", "description": "段落标题"},
                        "content": {"type": "string", "description": "段落内容"},
                    },
                },
                "description": "文档各章节内容",
            },
            "output_path": {"type": "string", "description": "输出文件路径"},
        },
        "required": ["format", "title", "sections", "output_path"],
    }
    is_concurrency_safe: bool = True
    is_read_only: bool = False

    async def check_permissions(
        self, tool_input: dict, context: PermissionContext,
    ) -> PermissionDecision:
        return PermissionDecision(
            behavior=PermissionBehavior.ASK,
            message="Document generation creates files.",
        )

    async def call(
        self,
        format: str,
        title: str,
        sections: List[dict],
        output_path: str,
    ) -> ToolChunk:
        try:
            if format == "docx":
                return await self._generate_docx(title, sections, output_path)
            elif format == "pdf":
                return await self._generate_pdf(title, sections, output_path)
            elif format == "pptx":
                return await self._generate_pptx(title, sections, output_path)
            else:
                return _text_chunk({"error": f"不支持的格式: {format}"})
        except Exception as exc:
            logger.warning("document_generator failed: %s", exc)
            return _text_chunk({"error": str(exc)})

    async def _generate_docx(self, title: str, sections: List[dict], output_path: str) -> ToolChunk:
        from docx import Document
        doc = Document()
        doc.add_heading(title, level=0)
        for sec in sections:
            heading = sec.get("heading", "")
            content = sec.get("content", "")
            if heading:
                doc.add_heading(heading, level=1)
            if content:
                doc.add_paragraph(content)
        os.makedirs(os.path.dirname(output_path) or ".", exist_ok=True)
        doc.save(output_path)
        return _text_chunk({
            "file_path": output_path,
            "format": "docx",
            "sections": len(sections),
            "size_bytes": os.path.getsize(output_path),
        })

    async def _generate_pdf(self, title: str, sections: List[dict], output_path: str) -> ToolChunk:
        from reportlab.lib.pagesizes import A4
        from reportlab.pdfgen import canvas
        from reportlab.lib.units import mm
        from reportlab.pdfbase import pdfmetrics
        from reportlab.pdfbase.ttfonts import TTFont

        os.makedirs(os.path.dirname(output_path) or ".", exist_ok=True)
        c = canvas.Canvas(output_path, pagesize=A4)
        width, height = A4

        # 尝试注册中文字体
        font_name = "Helvetica-Bold"
        font_name_body = "Helvetica"
        for font_path in [
            "C:/Windows/Fonts/simhei.ttf",
            "C:/Windows/Fonts/msyh.ttf",
            "/usr/share/fonts/truetype/wqy/wqy-zenhei.ttc",
        ]:
            if os.path.exists(font_path):
                try:
                    pdfmetrics.registerFont(TTFont("ChineseFont", font_path))
                    font_name = "ChineseFont"
                    font_name_body = "ChineseFont"
                    break
                except Exception:
                    continue

        y = height - 40 * mm
        c.setFont(font_name, 18)
        c.drawString(30 * mm, y, title)
        y -= 15 * mm

        for sec in sections:
            heading = sec.get("heading", "")
            content = sec.get("content", "")
            if y < 30 * mm:
                c.showPage()
                y = height - 30 * mm
            if heading:
                c.setFont(font_name, 14)
                c.drawString(30 * mm, y, heading)
                y -= 8 * mm
            if content:
                c.setFont(font_name_body, 11)
                # 简单换行处理
                max_chars = 60
                lines = [content[i:i + max_chars] for i in range(0, len(content), max_chars)]
                for line in lines:
                    if y < 20 * mm:
                        c.showPage()
                        y = height - 30 * mm
                    c.drawString(30 * mm, y, line)
                    y -= 6 * mm
            y -= 4 * mm

        c.save()
        return _text_chunk({
            "file_path": output_path,
            "format": "pdf",
            "sections": len(sections),
            "size_bytes": os.path.getsize(output_path),
        })

    async def _generate_pptx(self, title: str, sections: List[dict], output_path: str) -> ToolChunk:
        from pptx import Presentation
        from pptx.util import Inches, Pt

        prs = Presentation()

        # 标题页
        slide = prs.slides.add_slide(prs.slide_layouts[0])
        slide.shapes.title.text = title
        if slide.placeholders[1]:
            slide.placeholders[1].text = ""

        # 内容页
        for sec in sections:
            slide = prs.slides.add_slide(prs.slide_layouts[1])
            heading = sec.get("heading", "")
            content = sec.get("content", "")
            if slide.shapes.title:
                slide.shapes.title.text = heading
            body_shape = slide.placeholders[1]
            if body_shape:
                tf = body_shape.text_frame
                tf.text = content

        os.makedirs(os.path.dirname(output_path) or ".", exist_ok=True)
        prs.save(output_path)
        return _text_chunk({
            "file_path": output_path,
            "format": "pptx",
            "slides": len(sections) + 1,
            "size_bytes": os.path.getsize(output_path),
        })


# ============= FormatConverter =============

class FormatConverter(ToolBase):
    """文档格式转换。"""

    name: str = "format_converter"
    description: str = "文档格式互转：PDF→Markdown、Markdown→Word、Word→PDF 等"
    input_schema: dict = {
        "type": "object",
        "properties": {
            "input_path": {"type": "string", "description": "输入文件路径"},
            "output_format": {
                "type": "string",
                "enum": ["md", "docx", "pdf"],
                "description": "目标格式",
            },
            "output_path": {
                "type": "string",
                "description": "输出文件路径（不传则自动生成）",
            },
        },
        "required": ["input_path", "output_format"],
    }
    is_concurrency_safe: bool = True
    is_read_only: bool = False

    async def check_permissions(
        self, tool_input: dict, context: PermissionContext,
    ) -> PermissionDecision:
        return PermissionDecision(
            behavior=PermissionBehavior.ASK,
            message="Format conversion creates new files.",
        )

    async def call(
        self,
        input_path: str,
        output_format: str,
        output_path: Optional[str] = None,
    ) -> ToolChunk:
        if not os.path.isfile(input_path):
            return _text_chunk({"error": f"输入文件不存在: {input_path}"})

        input_ext = Path(input_path).suffix.lower()
        if not output_path:
            output_path = str(Path(input_path).with_suffix(f".{output_format}"))

        try:
            if input_ext == ".pdf" and output_format == "md":
                return await self._pdf_to_md(input_path, output_path)
            elif input_ext == ".md" and output_format == "docx":
                return await self._md_to_docx(input_path, output_path)
            elif input_ext == ".md" and output_format == "pdf":
                return await self._md_to_pdf(input_path, output_path)
            elif input_ext == ".docx" and output_format == "pdf":
                return await self._docx_to_pdf(input_path, output_path)
            elif input_ext == ".docx" and output_format == "md":
                return await self._docx_to_md(input_path, output_path)
            else:
                return _text_chunk({
                    "error": f"不支持的转换: {input_ext} → {output_format}",
                    "supported": [
                        "pdf → md", "md → docx", "md → pdf",
                        "docx → pdf", "docx → md",
                    ],
                })
        except Exception as exc:
            logger.warning("format_converter failed: %s", exc)
            return _text_chunk({"error": str(exc)})

    async def _pdf_to_md(self, input_path: str, output_path: str) -> ToolChunk:
        import pdfplumber
        lines = []
        with pdfplumber.open(input_path) as pdf:
            for i, page in enumerate(pdf.pages):
                text = page.extract_text() or ""
                if text.strip():
                    lines.append(f"<!-- Page {i + 1} -->\n\n{text}")
                tables = page.extract_tables() or []
                for table in tables:
                    if table:
                        header = table[0]
                        lines.append("| " + " | ".join(str(c or "") for c in header) + " |")
                        lines.append("| " + " | ".join("---" for _ in header) + " |")
                        for row in table[1:]:
                            lines.append("| " + " | ".join(str(c or "") for c in row) + " |")
                        lines.append("")

        os.makedirs(os.path.dirname(output_path) or ".", exist_ok=True)
        with open(output_path, "w", encoding="utf-8") as f:
            f.write("\n\n".join(lines))
        return _text_chunk({
            "output_path": output_path,
            "format": "md",
            "size_bytes": os.path.getsize(output_path),
        })

    async def _md_to_docx(self, input_path: str, output_path: str) -> ToolChunk:
        from docx import Document
        with open(input_path, "r", encoding="utf-8") as f:
            content = f.read()

        doc = Document()
        for line in content.split("\n"):
            stripped = line.strip()
            if stripped.startswith("# "):
                doc.add_heading(stripped[2:], level=1)
            elif stripped.startswith("## "):
                doc.add_heading(stripped[3:], level=2)
            elif stripped.startswith("### "):
                doc.add_heading(stripped[4:], level=3)
            elif stripped:
                doc.add_paragraph(stripped)

        os.makedirs(os.path.dirname(output_path) or ".", exist_ok=True)
        doc.save(output_path)
        return _text_chunk({
            "output_path": output_path,
            "format": "docx",
            "size_bytes": os.path.getsize(output_path),
        })

    async def _md_to_pdf(self, input_path: str, output_path: str) -> ToolChunk:
        import markdown
        from reportlab.lib.pagesizes import A4
        from reportlab.pdfgen import canvas
        from reportlab.lib.units import mm

        with open(input_path, "r", encoding="utf-8") as f:
            md_content = f.read()

        # 简单 Markdown → PDF（按行处理）
        os.makedirs(os.path.dirname(output_path) or ".", exist_ok=True)
        c = canvas.Canvas(output_path, pagesize=A4)
        width, height = A4
        y = height - 25 * mm

        for line in md_content.split("\n"):
            if y < 20 * mm:
                c.showPage()
                y = height - 25 * mm
            stripped = line.strip()
            if stripped.startswith("# "):
                c.setFont("Helvetica-Bold", 18)
                c.drawString(25 * mm, y, stripped[2:])
                y -= 10 * mm
            elif stripped.startswith("## "):
                c.setFont("Helvetica-Bold", 14)
                c.drawString(25 * mm, y, stripped[3:])
                y -= 8 * mm
            elif stripped:
                c.setFont("Helvetica", 11)
                c.drawString(25 * mm, y, stripped[:100])
                y -= 6 * mm
            else:
                y -= 3 * mm

        c.save()
        return _text_chunk({
            "output_path": output_path,
            "format": "pdf",
            "size_bytes": os.path.getsize(output_path),
        })

    async def _docx_to_pdf(self, input_path: str, output_path: str) -> ToolChunk:
        from docx import Document
        from reportlab.lib.pagesizes import A4
        from reportlab.pdfgen import canvas
        from reportlab.lib.units import mm

        doc = Document(input_path)
        os.makedirs(os.path.dirname(output_path) or ".", exist_ok=True)
        c = canvas.Canvas(output_path, pagesize=A4)
        width, height = A4
        y = height - 25 * mm

        for para in doc.paragraphs:
            if y < 20 * mm:
                c.showPage()
                y = height - 25 * mm
            text = para.text.strip()
            if not text:
                y -= 3 * mm
                continue
            if para.style and para.style.name and "Heading" in para.style.name:
                c.setFont("Helvetica-Bold", 14)
            else:
                c.setFont("Helvetica", 11)
            c.drawString(25 * mm, y, text[:100])
            y -= 6 * mm

        c.save()
        return _text_chunk({
            "output_path": output_path,
            "format": "pdf",
            "size_bytes": os.path.getsize(output_path),
        })

    async def _docx_to_md(self, input_path: str, output_path: str) -> ToolChunk:
        from docx import Document
        doc = Document(input_path)
        lines = []
        for para in doc.paragraphs:
            text = para.text.strip()
            if not text:
                lines.append("")
                continue
            if para.style and para.style.name:
                style = para.style.name
                if style == "Heading 1":
                    lines.append(f"# {text}")
                elif style == "Heading 2":
                    lines.append(f"## {text}")
                elif style == "Heading 3":
                    lines.append(f"### {text}")
                else:
                    lines.append(text)
            else:
                lines.append(text)

        os.makedirs(os.path.dirname(output_path) or ".", exist_ok=True)
        with open(output_path, "w", encoding="utf-8") as f:
            f.write("\n".join(lines))
        return _text_chunk({
            "output_path": output_path,
            "format": "md",
            "size_bytes": os.path.getsize(output_path),
        })
